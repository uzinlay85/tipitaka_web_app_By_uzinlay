"""Export Tipitaka book pages to Word (.docx).

The page HTML stored in the DB (classes: nikaya/book/chapter/bodytext/...)
is parsed into plain blocks, then rendered into a .docx via python-docx
(pure python, no system deps; Word shapes Myanmar correctly).
"""
import io
import re
from html.parser import HTMLParser

# p class -> block kind
_BLOCK_KIND = {
    "nikaya": "nikaya",
    "book": "book",
    "title": "book",
    "chapter": "chapter",
    "subhead": "chapter",
    "sub-head": "chapter",
    "centered": "center",
    "center": "center",
    "bodytext": "body",
    "noindentbodytext": "body",
    "body": "body",
    "hangnum": "body",
    "gatha": "verse",
    "gatha1": "verse",
    "gatha2": "verse",
    "gatha3": "verse",
}


class _BlockParser(HTMLParser):
    """Turn page HTML into a list of (kind, segments).

    segments: list of (text, bold) tuples; paranum spans become bold.
    """

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks = []
        self._kind = None
        self._segs = []
        self._bold = False
        self._skip_depth = 0  # inside <a ...> anchors (they are empty markers)

    def _flush(self):
        if self._kind is None:
            self._segs = []
            self._bold = False
            return
        text = "".join(t for t, _ in self._segs)
        text = re.sub(r"[ \t\xa0]+", " ", text).strip()
        # drop blocks that are only anchors/whitespace
        if text:
            # normalize internal newlines from <br>
            text = re.sub(r"\n{3,}", "\n\n", text)
            segs = [(t, b) for t, b in self._segs if t.strip() or t in (" ", "\n")]
            # strip leading/trailing word-space segs
            while segs and segs[0][0] == " ":
                segs.pop(0)
            while segs and segs[-1][0] == " ":
                segs.pop()
            if segs:
                self.blocks.append((self._kind, segs))
        self._kind = None
        self._segs = []
        self._bold = False

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == "a":
            self._skip_depth += 1
            return
        if self._skip_depth:
            return
        if tag == "p":
            self._flush()
            cls = (attrs.get("class") or "").split()
            kind = "body"
            for c in cls:
                if c in _BLOCK_KIND:
                    kind = _BLOCK_KIND[c]
                    break
            self._kind = kind
            self._segs = []
        elif tag == "br":
            if self._kind:
                self._segs.append(("\n", False))
        elif tag == "span" and "paranum" in (attrs.get("class") or ""):
            self._bold = True

    def handle_endtag(self, tag):
        if tag == "a":
            self._skip_depth = max(0, self._skip_depth - 1)
            return
        if self._skip_depth:
            return
        if tag == "p":
            self._flush()
        elif tag == "span":
            self._bold = False

    def handle_data(self, data):
        if self._skip_depth or self._kind is None:
            return
        if not data.strip():
            # Whitespace between inline spans (e.g. </span> <span> from the
            # web reader's pali-word/no-split/gatha-pada wrappers) is a real
            # word space - keep one, avoid leading/dupes.
            if self._segs and not self._segs[-1][0].endswith((" ", "\n")):
                self._segs.append((" ", False))
            return
        self._segs.append((data, self._bold))


def parse_page(html_text):
    """Parse one page's HTML -> [(kind, [(text, bold), ...]), ...]."""
    p = _BlockParser()
    p.feed(html_text or "")
    p._flush()
    return p.blocks


def _iter_all_blocks(pages):
    """pages: list of (page_num, blocks)."""
    for page_num, blocks in pages:
        yield page_num, blocks


# ---------------------------------------------------------------- DOCX ---

def build_docx(book_name, pages, edition_label=""):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Base style: Pyidaungsu for Myanmar shaping in Word
    normal = doc.styles["Normal"]
    normal.font.name = "Pyidaungsu"
    normal.font.size = Pt(13)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.5

    for style_name, size, bold in (
        ("Heading 1", 17, True),
        ("Heading 2", 15, True),
        ("Heading 3", 13, True),
    ):
        st = doc.styles[style_name]
        st.font.name = "Pyidaungsu"
        st.font.size = Pt(size)
        st.font.bold = bold

    # Title page header
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(book_name)
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.name = "Pyidaungsu"
    if edition_label:
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = sub.add_run(edition_label)
        r.font.size = Pt(13)
        r.font.name = "Pyidaungsu"
        r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    doc.add_page_break()

    total = len(pages)
    for idx, (page_num, blocks) in enumerate(_iter_all_blocks(pages)):
        label = doc.add_paragraph()
        label.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = label.add_run("— စာမျက်နှာ %d —" % page_num)
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        r.font.name = "Pyidaungsu"

        for kind, segs in blocks:
            if kind == "nikaya":
                par = doc.add_paragraph(style="Heading 1")
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif kind == "book":
                par = doc.add_paragraph(style="Heading 1")
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif kind == "chapter":
                par = doc.add_paragraph(style="Heading 2")
            elif kind == "center":
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif kind == "verse":
                par = doc.add_paragraph()
                par.paragraph_format.left_indent = Pt(24)
            else:
                par = doc.add_paragraph()
            for text, bold in segs:
                # split on <br> newlines into separate runs
                for j, chunk in enumerate(text.split("\n")):
                    if j:
                        par.add_run().add_break()
                    run = par.add_run(chunk)
                    run.font.name = "Pyidaungsu"
                    if bold:
                        run.font.bold = True
        if idx < total - 1:
            doc.add_page_break()

    core = doc.core_properties
    core.title = book_name
    core.author = "Tipitaka Web App"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


MAX_EXPORT_PAGES = 1000
